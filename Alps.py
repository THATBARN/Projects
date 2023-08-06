import os
import requests
import tempfile
import time

from datetime import datetime
from dotenv import dotenv_values
from openpyxl import load_workbook

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait


env_conf = dotenv_values(os.path.abspath(os.path.dirname(__file__)) + "/.env")
print("Get new market requests")

# GET new order ids
url = "https://admin.bozoraka.com/api/new-market-request-ids?sellerId=25"
headers = {
    'Content-type': 'application/json',
    'Authorization': env_conf.get("BOZORAKA_AUTH")
}
response = requests.get(url, headers=headers)
order_ids = response.text

# POST order ids to receive order excel file
url = "https://admin.bozoraka.com/api/new-market-requests-excel"
response = requests.post(url, data=order_ids, headers=headers)

new_requests = tempfile.NamedTemporaryFile(prefix="alps_")
new_requests.write(response.content)

# Remove multi-box orders
wb = load_workbook(filename=new_requests)
sheet = wb.active

# Put order numbers in list
order_numbers = []
columnA = sheet["A"]
for cell in columnA:
    order_numbers.append(cell.value)
order_numbers.remove(None)

# Cycle through list to find dupes
dupe_orders = []
for order in order_numbers:
    if order.count("-") > 2:
        dupe_orders.append(order)
        original_order = order[0:-2]
        if original_order not in dupe_orders:
            dupe_orders.append(original_order)

# Delete dupe rows
row = 1
while row <= sheet.max_row:
    if sheet.cell(row=row, column=1).value in dupe_orders:
        sheet.delete_rows(row, 1)
    else:
        row +=1

now = datetime.now()
orders_xlsx = now.strftime("%Y%m%d-%H%M") + ".xlsx"
wb.save(env_conf.get("ORDER_FOLDER_PATH") + orders_xlsx)

print("Uploading to ALPS website ...")
options = Options()
options.headless = True
driver = webdriver.Chrome(env_conf.get("CHROME_DRIVER_PATH"), options=options)

driver.get(env_conf.get("ALPS_URL"))
driver.find_element(
    By.XPATH, '//*[@id="principal"]/input').send_keys(env_conf.get("ALPS_ID"))
driver.find_element(
    By.XPATH, '//*[@id="credential"]/input').send_keys(env_conf.get("ALPS_PASS"))

login = driver.find_element(By.XPATH, '//*[@id="btn-login"]')
login.click()

time.sleep(3)

driver.find_element(
    By.XPATH, '/html/body/div[3]/header/nav/div[1]/div[1]/i-button'
).click()

driver.find_element(
    By.XPATH, '/html/body/div[3]/header/nav/div[3]/div/div[2]/div[7]/div[2]/ul/li/a'
).click()

time.sleep(3)

workframe = driver.find_element(By.ID, "workframe_10782")
driver.switch_to.frame(workframe)
driver.execute_script("arguments[0].removeAttribute('readonly')",
                      WebDriverWait(driver, 20).until(EC.element_to_be_clickable((By.XPATH, '//*[@id="cboUsrFmat01"]/input[2]'))))

typeinput = driver.find_element(By.XPATH, '//*[@id="cboUsrFmat01"]/input[2]')
typeinput.send_keys(Keys.ARROW_DOWN, Keys.ARROW_DOWN,
                    Keys.ARROW_DOWN, Keys.ENTER)

file_upload = driver.find_element(
    By.XPATH, '//*[@id="files"]'
)

file_upload.send_keys(env_conf.get("ORDER_FOLDER_PATH") + orders_xlsx)
time.sleep(30) # TODO: check the actual page response

print("Done uploading new market requests")
